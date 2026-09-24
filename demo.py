# -*- coding: utf-8 -*-
"""End-to-end demo on a synthetic two-column CV (no personal data).

1. Build a small CV with python-docx: a sidebar and a main column in one table.
2. Refuse sidebar entries that would wrap, using real font metrics.
3. Grow the skills list in place with cvlib, save a tailored copy.
4. Render both with LibreOffice and check that no sidebar section is split.

Requires LibreOffice (soffice), poppler (pdfinfo, pdftotext) and the Lato font.
"""
from docx import Document
from docx.shared import Pt, Inches
import cvlib as C

SKILLS = ["Product strategy", "Customer discovery", "API integrations", "Pricing & packaging"]
MORE = SKILLS + ["Generative AI & LLMs", "Stakeholder alignment", "Contract negotiation"]
LANGS = ["English - C2", "German - C2", "French - C1"]


def build(path):
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.7)
    style = doc.styles['Normal']
    style.font.name = 'Lato'
    style.font.size = Pt(11)
    t = doc.add_table(rows=2, cols=2)
    t.autofit = False
    widths = (Inches(2.1), Inches(4.9))
    for row in t.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w
    t.cell(0, 1).paragraphs[0].add_run('Jane Example').font.size = Pt(26)
    side, main = t.cell(1, 0), t.cell(1, 1)
    side.paragraphs[0].add_run('SKILLS').bold = True
    for s in SKILLS:
        side.add_paragraph(s)
    side.add_paragraph().add_run('LANGUAGES').bold = True
    for s in LANGS:
        side.add_paragraph(s)
    main.paragraphs[0].add_run('EXPERIENCE').bold = True
    for i in range(6):
        main.add_paragraph('Head of Product, Example AG. Led a cross-functional team that took a '
                           'voice product from pilot to production, owned the roadmap and the budget, '
                           'and worked with sales on enterprise deals across Europe.')
    doc.save(path)


if __name__ == '__main__':
    build('sample.docx')
    C.assert_one_line(MORE)            # fails loudly if any entry would wrap
    print('all sidebar entries fit on one line')

    root = C.load('sample.docx')
    ps = C.paras(root)
    idx = [i for i, p in enumerate(ps) if C.text_of(p) in SKILLS]
    C.fill_list(C.grab(root, idx), MORE)   # capture first, then mutate
    C.save('sample.docx', 'tailored.docx', root)

    for f in ('sample.docx', 'tailored.docx'):
        pdf, pages = C.page_texts(f)
        split = C.sections_intact(pdf, {'SKILLS': MORE[-1] if f == 'tailored.docx' else SKILLS[-1],
                                        'LANGUAGES': LANGS[-1]})
        print(f'{f}: pages={len(pages)}  split sections: {split or "none"}')
